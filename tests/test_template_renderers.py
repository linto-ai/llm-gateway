#!/usr/bin/env python3
"""Template renderers (app/services/template_renderers): one test class per renderer, run through
DocumentService.substitute_placeholders / ExportService._get_missing_placeholders like in production."""
from pathlib import Path

import pytest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from app.services import template_renderers
from app.services.document_service import DocumentService
from app.services.document_template_service import DocumentTemplateService
from app.services.export_service import ExportService
from app.services.template_renderers import hidden_sections

LISTS = template_renderers.LISTS_KEY


def _render(doc, values, props=None):
    values.setdefault("output", "")
    with patch("app.services.template_renderers.template_properties", lambda d: dict(props or {})):
        DocumentService.__new__(DocumentService).substitute_placeholders(doc, values)
    return doc


class TestRepeatedRows:
    def _doc(self):
        doc = Document()
        t = doc.add_table(rows=2, cols=3)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(0, 2).text = "Porteur", "Action", "Échéance"
        t.cell(1, 0).paragraphs[0].add_run("{{actions.porteur}}").bold = True
        t.cell(1, 1).paragraphs[0].add_run("{{actions: les actions}}{{actions.action: 12 mots max}}")
        t.cell(1, 2).paragraphs[0].add_run("{{actions.echeance}}")
        return doc

    def test_row_repeated_per_item_keeping_format(self):
        items = [{"porteur": "Claire", "action": "Envoyer le devis", "echeance": "vendredi"},
                 {"porteur": "Hugo", "action": "Réserver la salle", "echeance": ""}]
        rows = _render(self._doc(), {LISTS: {"actions": items}}).tables[0].rows
        assert [c.text for c in rows[1].cells] == ["Claire", "Envoyer le devis", "vendredi"]
        assert [c.text for c in rows[2].cells] == ["Hugo", "Réserver la salle", ""]
        assert rows[2].cells[0].paragraphs[0].runs[0].bold is True

    def test_empty_list_removes_row_and_missing_list_leaves_blank_row(self):
        assert len(_render(self._doc(), {LISTS: {"actions": []}}).tables[0].rows) == 1
        rows = _render(self._doc(), {LISTS: {}}).tables[0].rows
        assert len(rows) == 2 and all(c.text == "" for c in rows[1].cells)

    def test_list_of_strings(self):
        doc = Document()
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{points.valeur}}"
        doc = _render(doc, {LISTS: {"points": ["un", "deux"]}})
        assert [r.cells[0].text for r in doc.tables[0].rows] == ["un", "deux"]

    def test_extraction_groups_dotted_fields_into_one_request(self):
        svc = ExportService.__new__(ExportService)
        svc.template_service = DocumentTemplateService.__new__(DocumentTemplateService)
        placeholders = ["titre: court", "actions: les actions", "actions.porteur",
                        "actions.action: 12 mots max", "output"]
        missing = svc._get_missing_placeholders(placeholders, {})
        lists = [m for m in missing if m.startswith("actions:")]
        assert "titre: court" in missing and len(lists) == 1
        assert "list of JSON objects" in lists[0] and "action (12 mots max)" in lists[0] and "les actions" in lists[0]
        assert svc._get_missing_placeholders(["actions.porteur"], {"actions": []}) == []


class TestRichValues:
    def test_multiline_value_becomes_paragraphs_with_bullets_and_bold(self):
        doc = Document()
        doc.add_paragraph("Intro {{x}}")
        doc.add_paragraph().add_run("{{decisions}}").italic = True
        doc = _render(doc, {"x": "a", "decisions": "• **Claire** : valider le devis\n- Hugo : relancer"})
        assert [p.text for p in doc.paragraphs] == ["Intro a", "• Claire : valider le devis", "• Hugo : relancer"]
        first = doc.paragraphs[1]
        assert [r.text for r in first.runs if r.bold] == ["Claire"]
        assert all(r.italic for r in first.runs) and first.paragraph_format.first_line_indent < 0

    def test_value_inside_sentence_untouched_and_cells_supported(self):
        doc = Document()
        doc.add_paragraph("Voir {{d}} ici")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "{{p}}"
        doc = _render(doc, {"d": "a\nb", "p": "• A\n• B"})
        assert len(doc.paragraphs) == 1
        assert [p.text for p in doc.tables[0].cell(0, 0).paragraphs] == ["• A", "• B"]


TABLE_MD = """## Sujets

| # | Sujet | Type | Porteur | Échéance |
|---|---|---|---|---|
| 01 | Premier point | O ▲ | - | - |
| 02 | Claire enverra le devis | A ► | Claire Martin | vendredi |
## Réunion
- Projet : Test
"""


class TestHiddenSections:
    def test_section_removed_case_insensitive_with_its_content(self):
        out = hidden_sections.prepare(TABLE_MD, {"hide_sections": "réunion"})
        assert "Projet : Test" not in out and "| 01 |" in out
        out = hidden_sections.prepare(TABLE_MD, {"hide_sections": "SUJETS"})
        assert "| 01 |" not in out and "## Réunion" in out

    def test_section_ends_at_next_heading_of_same_level(self):
        md = "## A\ntexte a\n### A1\nsous a\n## B\ntexte b"
        assert hidden_sections.prepare(md, {"hide_sections": "A"}) == "## B\ntexte b"

    def test_blank_line_forced_after_table(self):
        assert "| vendredi |\n\n## Réunion" in hidden_sections.prepare(TABLE_MD, {})
        doc = Document()
        doc.add_paragraph("{{output}}")
        doc = _render(doc, {"output": TABLE_MD})
        assert len(doc.tables[0].rows) == 3


class TestOutputTables:
    def _doc(self, style=None):
        doc = Document()
        if style:
            doc.styles.add_style(style, WD_STYLE_TYPE.TABLE)
        doc.add_paragraph("{{output}}")
        return doc

    def test_style_widths_align_and_empty_cells(self):
        props = {"style_table": "CRSujets", "table_widths": "5,55,7,20,13",
                 "table_align": "center,left,center,center,center"}
        doc = _render(self._doc("CRSujets"), {"output": TABLE_MD}, props)
        tbl = doc.tables[0]._tbl
        assert tbl.tblPr.find(qn("w:tblStyle")).get(qn("w:val")) == "CRSujets"
        widths = [int(g.get(qn("w:w"))) for g in tbl.tblGrid.findall(qn("w:gridCol"))]
        assert widths[1] > 10 * widths[0]
        assert all(tr.trPr.find(qn("w:cantSplit")) is not None for tr in tbl.findall(qn("w:tr")))
        assert list(tbl.findall(qn("w:tr"))[0].iter(qn("w:b"))) == []
        row = doc.tables[0].rows[1]
        assert row.cells[3].text == "" and row.cells[2].paragraphs[0]._p.pPr.jc.get(qn("w:val")) == "center"

    def test_unknown_style_ignored_and_no_property_no_change(self):
        doc = _render(self._doc(), {"output": TABLE_MD}, {"style_table": "Nope"})
        assert doc.tables[0]._tbl.tblPr.find(qn("w:tblStyle")) is None
        assert doc.tables[0].rows[1].cells[3].text == "-"
        doc = _render(self._doc(), {"output": TABLE_MD})
        assert doc.tables[0]._tbl.tblPr.find(qn("w:tblLayout")) is None

    def test_symbol_colors(self):
        doc = _render(self._doc(), {"output": TABLE_MD}, {"table_symbol_colors": "►:C51C42,▲:2B5797"})
        runs = doc.tables[0].rows[2].cells[2].paragraphs[0].runs
        assert "".join(r.text for r in runs) == "A ►"
        colored = [r for r in runs if r.font.color and r.font.color.rgb is not None]
        assert [(r.text, str(r.font.color.rgb)) for r in colored] == [("►", "C51C42")]


def test_docs_in_sync_with_specs():
    doc = Path(__file__).resolve().parent.parent / "docs" / "TEMPLATE_RENDERERS.md"
    if not doc.exists():
        pytest.skip("docs/ not available (container mounts app/ and tests/ only)")
    assert doc.read_text(encoding="utf-8") == template_renderers.render_docs() + "\n", \
        "run: python -c 'from app.services.template_renderers import render_docs; print(render_docs())' > docs/TEMPLATE_RENDERERS.md"


class TestPluginRegistry:
    def test_plugins_discovered_with_their_family(self):
        from app.services.template_renderers.base import REGISTRY
        assert {(r.name, r.family) for r in REGISTRY} == {
            ("repeated_rows", "placeholder"), ("rich_values", "placeholder"),
            ("hidden_sections", "output"), ("output_tables", "output"), ("mindmap", "placeholder")}

    def test_register_rejects_invalid_plugins(self):
        import pytest as _pytest
        from app.services.template_renderers.base import OUTPUT, Renderer, RendererSpec, register
        spec = RendererSpec("t", "s", "p", "e")

        class Foreign(Renderer):
            name, family, spec_ = "foreign", OUTPUT, None
            def before_substitution(self, doc, ctx):
                pass
        Foreign.spec = spec
        with _pytest.raises(TypeError, match="another family"):
            register(Foreign)

        class NoSpec(Renderer):
            name, family = "nospec", OUTPUT
        with _pytest.raises(TypeError, match="required"):
            register(NoSpec)

        class Duplicate(Renderer):
            name, family = "output_tables", OUTPUT
        Duplicate.spec = spec
        with _pytest.raises(TypeError, match="already registered"):
            register(Duplicate)

    def test_new_plugin_runs_at_its_hook_in_order(self):
        from app.services.template_renderers.base import OUTPUT, REGISTRY, Renderer, RendererSpec, register

        @register
        class Shout(Renderer):
            name, family, order = "test_shout", OUTPUT, 99
            spec = RendererSpec("t", "s", "p", "e")
            def prepare_output(self, doc, markdown, ctx):
                return markdown.upper()
        try:
            out = template_renderers.prepare_output(Document(), "| a |\n## Réunion\n- x")
            assert out == "| A |\n\n## RÉUNION\n- X"  # hidden_sections (order 10) ran first
        finally:
            REGISTRY[:] = [r for r in REGISTRY if r.name != "test_shout"]


class TestMindmap:
    OUTLINE = "Projet Atlas\n- Budget\n  - **4 200 €** de licences\n  - Renouvellement en 2028\n- Infrastructure\n  - Serveurs dédiés (80 %)\n- Organisation"

    def test_outline_parsing(self):
        from app.services.template_renderers.mindmap import parse_outline
        center, branches = parse_outline(self.OUTLINE)
        assert center == "Projet Atlas"
        assert branches == [("Budget", ["4 200 € de licences", "Renouvellement en 2028"]),
                            ("Infrastructure", ["Serveurs dédiés (80 %)"]), ("Organisation", [])]

    def test_placeholder_becomes_image(self):
        doc = Document()
        doc.add_paragraph("Avant")
        doc.add_paragraph("{{mindmap_sujets: les sujets}}")
        doc = _render(doc, {"mindmap_sujets": self.OUTLINE})
        para = doc.paragraphs[1]
        assert para.text == ""
        assert para._p.findall(".//" + qn("w:drawing"))
        assert doc.inline_shapes[0].width > 0

    def test_not_extracted_leaves_empty_paragraph_and_bad_outline_falls_back_to_text(self):
        doc = Document()
        doc.add_paragraph("{{mindmap_sujets}}")
        doc.add_paragraph("{{mindmap_autre}}")
        doc = _render(doc, {"mindmap_autre": "juste une phrase"})
        assert doc.paragraphs[0].text == "" and not doc.inline_shapes
        assert doc.paragraphs[1].text == "juste une phrase"

    def test_extraction_request_is_an_outline(self):
        svc = ExportService.__new__(ExportService)
        svc.template_service = DocumentTemplateService.__new__(DocumentTemplateService)
        missing = svc._get_missing_placeholders(["mindmap_sujets: les sujets abordés", "titre"], {})
        req = [m for m in missing if m.startswith("mindmap_sujets:")]
        assert len(req) == 1 and "outline" in req[0] and "les sujets abordés" in req[0]
        assert "titre" in missing


def test_end_of_job_extraction_fields_use_renderers():
    from app.services.template_renderers import prepare_extraction_fields
    fields = ["titre: court", "actions: les actions", "actions.porteur", "actions.action: 12 mots",
              "mindmap_sujets: les sujets", "output"]
    out = prepare_extraction_fields(fields)
    assert "titre: court" in out and "output" in out
    assert not any(f.startswith("actions.") for f in out)
    assert sum(f.startswith("actions:") for f in out) == 1 and "list of JSON objects" in [f for f in out if f.startswith("actions:")][0]
    assert [f for f in out if f.startswith("mindmap_sujets:")][0].count("outline") == 1
