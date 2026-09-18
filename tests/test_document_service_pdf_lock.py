#!/usr/bin/env python3
"""Tests for the footer note and the locked PDF export of DocumentService."""
import re
import shutil
from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION

from app.services.document_service import DocumentService

NOTE = "Généré avec LinTO Studio · linto.ai"

# PDF permission bits (ISO 32000-1, table 22)
PRINT = 1 << 2
MODIFY = 1 << 3
COPY = 1 << 4

needs_libreoffice = pytest.mark.skipif(
    shutil.which("libreoffice") is None, reason="LibreOffice not installed"
)


def _footer_texts(footer):
    return [p.text for p in footer.paragraphs]


def _docx_buffer(text="Compte rendu"):
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _permissions(pdf: bytes):
    """/P of the Encrypt dictionary, None when the PDF is not encrypted."""
    if b"/Encrypt" not in pdf:
        return None
    match = re.search(rb"/P\s*(-?\d+)", pdf)
    assert match, "encrypted PDF without /P"
    return int(match.group(1))


class TestFooterNote:
    def test_added_once_to_the_first_footer(self):
        doc = Document()
        DocumentService._add_footer_note(doc, NOTE)
        assert _footer_texts(doc.sections[0].footer).count(NOTE) == 1

    def test_linked_sections_are_not_duplicated(self):
        doc = Document()
        doc.add_section(WD_SECTION.NEW_PAGE)
        DocumentService._add_footer_note(doc, NOTE)
        assert doc.sections[1].footer.is_linked_to_previous
        assert _footer_texts(doc.sections[0].footer).count(NOTE) == 1

    def test_unlinked_section_and_first_page_footers_get_it(self):
        doc = Document()
        doc.sections[0].different_first_page_header_footer = True
        second = doc.add_section(WD_SECTION.NEW_PAGE)
        second.footer.is_linked_to_previous = False
        DocumentService._add_footer_note(doc, NOTE)
        assert NOTE in _footer_texts(doc.sections[0].footer)
        assert NOTE in _footer_texts(doc.sections[0].first_page_footer)
        assert NOTE in _footer_texts(second.footer)

    def test_keeps_existing_footer_content(self):
        doc = Document()
        doc.sections[0].footer.paragraphs[0].text = "Page"
        DocumentService._add_footer_note(doc, NOTE)
        assert _footer_texts(doc.sections[0].footer) == ["Page", NOTE]


@needs_libreoffice
class TestLockedPdf:
    async def test_lock_forbids_editing_and_copying_but_not_printing(self):
        pdf = (await DocumentService()._convert_docx_to_pdf(_docx_buffer(), None, lock=True)).getvalue()
        perms = _permissions(pdf)
        assert perms is not None
        assert perms & PRINT
        assert not perms & MODIFY
        assert not perms & COPY

    async def test_no_lock_by_default(self):
        pdf = (await DocumentService()._convert_docx_to_pdf(_docx_buffer(), None)).getvalue()
        assert _permissions(pdf) is None

    def test_permission_password_is_never_reused(self):
        assert DocumentService._locked_pdf_filter() != DocumentService._locked_pdf_filter()
