"""Rich values: a placeholder alone in its paragraph or cell is rendered with formatting.

A multi-line value gives one paragraph per line; lines starting with •, -, * or 1. get a hanging
bullet indent; **text** becomes bold. Font, size and colour come from the placeholder run in the
template. A placeholder inside a sentence stays plain text.
"""
import re
from copy import deepcopy
from typing import Any, Dict

from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from .base import PLACEHOLDER, Renderer, RendererSpec, register

SPEC = RendererSpec(
    trigger="Un placeholder seul dans son paragraphe ou sa cellule, dont la valeur a plusieurs lignes ou du **gras**.",
    service_prompt="Aucune exigence propre.",
    placeholder_prompt=(
        "Demander explicitement la forme : « un élément par ligne commençant par • suivi d'un espace », "
        "« nom du porteur en gras ». Le prompt d'extraction autorise **…** et le saut de ligne \\n dans les valeurs."
    ),
    example="{{decisions_cles: les 4 décisions principales, une par ligne commençant par • suivi d'un espace, qui a tranché en gras}}",
)

BULLET_RE = re.compile(r"^\s*(?:[•\-*]|\d+[.)])\s+")


def _paragraphs(doc):
    seen = set()
    for para in doc.paragraphs:
        seen.add(id(para._p))
        yield para
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    if id(para._p) not in seen:  # merged cells repeat the same paragraphs
                        seen.add(id(para._p))
                        yield para


def enrich(doc, placeholders: Dict[str, Any]) -> None:
    values = {
        str(v).strip() for k, v in placeholders.items()
        if k != "output" and not k.startswith("__") and isinstance(v, str) and ("\n" in v.strip() or "**" in v)
    }
    if not values:
        return
    for para in list(_paragraphs(doc)):
        text = para.text.strip()
        if not text or text not in values:
            continue
        runs = [r for r in para.runs if r.text]
        if not runs:
            continue
        model_rpr = runs[0]._r.rPr
        anchor = para._p
        for i, line in enumerate(l for l in text.split("\n") if l.strip()):
            if i == 0:
                target = para
                for r in list(para.runs):
                    if r.text or r._r.find(qn("w:br")) is not None:
                        r._r.getparent().remove(r._r)
            else:
                new_p = deepcopy(para._p)
                for r_el in list(new_p.findall(qn("w:r"))):
                    new_p.remove(r_el)
                anchor.addnext(new_p)
                anchor = new_p
                target = Paragraph(new_p, para._parent)
            if BULLET_RE.match(line):
                line = BULLET_RE.sub("• ", line, count=1)
                target.paragraph_format.left_indent = Cm(0.35)
                target.paragraph_format.first_line_indent = Cm(-0.35)
            for j, chunk in enumerate(re.split(r"\*\*", line)):
                if not chunk:
                    continue
                run = target.add_run(chunk)
                if model_rpr is not None:
                    run._r.insert(0, deepcopy(model_rpr))
                if j % 2 == 1:
                    run.bold = True


@register
class RichValues(Renderer):
    """Rich values: a standalone multi-line or **bold** value rendered as formatted paragraphs."""
    name = "rich_values"
    family = PLACEHOLDER
    order = 90
    spec = SPEC

    def after_substitution(self, doc, ctx):
        enrich(doc, ctx.placeholders)
