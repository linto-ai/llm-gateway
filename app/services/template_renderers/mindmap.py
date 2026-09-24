"""Mindmap: a {{mindmap_...}} placeholder alone in its paragraph becomes a mind map image.

Extraction: the placeholder is requested as an indented outline string (central topic, then
"- branch" lines, then "  - detail" lines), a format LLMs produce far more reliably than nested JSON.
Rendering: central node, coloured branches split left and right, details next to their branch,
curved connectors; PNG drawn with Pillow and DejaVu Sans (present in the gateway image), inserted at
the width of the page body or of the table cell.
"""
import io
import re
from typing import Any, Dict, List, Tuple

from .base import PLACEHOLDER, Renderer, RendererSpec, logger, register

SPEC = RendererSpec(
    trigger="Placeholder dont le nom commence par mindmap_, seul dans son paragraphe ou sa cellule, par exemple {{mindmap_sujets: consigne}}.",
    service_prompt="Aucune exigence propre : la sortie du service doit contenir les thèmes et les points à cartographier (sections Sujets abordés et Discussion par exemple).",
    placeholder_prompt=(
        "La consigne dit quoi cartographier (thème central, branches, détails). Le gateway la transforme en demande "
        "de plan indenté : ligne 1 le thème central, puis « - branche », puis « - détail » indenté de deux espaces."
    ),
    example="{{mindmap_sujets: thème central = objet de la réunion ; une branche par sujet de la section Sujets abordés ; détails = faits clés de la section Discussion}}",
)

PREFIX = "mindmap_"
MAX_BRANCHES, MAX_LEAVES = 8, 5
PALETTE = ["1DAF92", "2B5797", "C47F00", "8E44AD", "C51C42", "0F7C8C", "5B6B2F", "B4589B"]
CENTER_FILL, INK, MUTED = "13806B", "2D2D2D", "6B7280"
FONT_DIR = "/usr/share/fonts/truetype/dejavu/"
PLACEHOLDER_RE = re.compile(r"\{\{\s*(" + PREFIX + r"[A-Za-z0-9_]*)\s*(?::[^}]*)?\}\}")


# ---------------------------------------------------------------- extraction

def outline_requests(placeholders: List[str], parse, current_metadata: Dict[str, Any], force: bool) -> Tuple[set, List[str]]:
    handled, requests = set(), []
    for placeholder in placeholders:
        info = parse(placeholder)
        name = info["name"]
        if not name.startswith(PREFIX) or name in handled:
            continue
        handled.add(name)
        if force or name not in current_metadata:
            hint = f" What to map: {info['description']}." if info["description"] else ""
            requests.append(
                f"{name}: an outline as ONE string, lines separated by \\n. Line 1: the central topic, 2 to 6 words. "
                f"Then each main branch on a line starting with '- ' (3 to {MAX_BRANCHES} branches, 2 to 5 words each), "
                f"each followed by its details on lines starting with two spaces and '- ' (1 to {MAX_LEAVES - 1} per branch, "
                f"8 words maximum each, facts and figures rather than generic words).{hint}"
            )
    return handled, requests


def parse_outline(text: str) -> Tuple[str, List[Tuple[str, List[str]]]]:
    center, branches = "", []
    for raw in str(text or "").replace("\r", "").split("\n"):
        if not raw.strip():
            continue
        indent = len(raw) - len(raw.lstrip(" \t").replace("\t", "  "))
        bullet = re.match(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$", raw)
        label = (bullet.group(1) if bullet else raw).replace("**", "").strip()
        if not label:
            continue
        if not bullet and not center and not branches:
            center = label
        elif bullet and indent < 2:
            branches.append((label, []))
        elif branches:
            branches[-1][1].append(label)
        elif not center:
            center = label
        else:
            branches.append((label, []))
    return center, [(b, leaves[:MAX_LEAVES]) for b, leaves in branches[:MAX_BRANCHES]]


# ---------------------------------------------------------------- drawing

def _font(size, bold=False):
    from PIL import ImageFont
    try:
        return ImageFont.truetype(FONT_DIR + ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"), size)
    except OSError:
        return ImageFont.load_default()


def _wrap(draw, text, font, max_w):
    # keep "80 %", "4 200 €", "Sujet :" together (French spacing before % € : ; ! ?)
    text = re.sub(r"(\d) (?=\d{3}\b)", "\\1\u00a0", text)
    text = re.sub(r" (?=[%€:;!?»])", "\u00a0", text)
    words, lines, cur = re.split(r"[ \t]+", text.strip()), [], ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if draw.textlength(trial, font=font) <= max_w or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines[:4]


def _rgb(hex_):
    return tuple(int(hex_[i:i + 2], 16) for i in (0, 2, 4))


def _curve(draw, p0, p1, color, width):
    (x0, y0), (x3, y3) = p0, p1
    mx = (x0 + x3) / 2
    pts = []
    for i in range(41):
        t = i / 40
        x = (1 - t) ** 3 * x0 + 3 * (1 - t) ** 2 * t * mx + 3 * (1 - t) * t ** 2 * mx + t ** 3 * x3
        y = (1 - t) ** 3 * y0 + 3 * (1 - t) ** 2 * t * y0 + 3 * (1 - t) * t ** 2 * y3 + t ** 3 * y3
        pts.append((x, y))
    draw.line(pts, fill=color, width=width, joint="curve")


def draw_mindmap(center: str, branches: List[Tuple[str, List[str]]]) -> bytes:
    from PIL import Image, ImageDraw

    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    f_center, f_branch, f_leaf = _font(46, True), _font(36, True), _font(33)
    margin, gap = 40, 34
    branch_w, leaf_w, branch_off, leaf_off = 380, 430, 90, 44
    c_lines = _wrap(probe, center, f_center, 420)
    c_w = (max(probe.textlength(l, font=f_center) for l in c_lines) + 70) if c_lines else 200
    # width from the content: centre + on each side branch box, connector and the widest detail text
    W = int(c_w + 2 * (branch_off + branch_w + leaf_off + 16 + leaf_w + margin))
    cx = W // 2

    right = branches[: (len(branches) + 1) // 2]
    left = branches[(len(branches) + 1) // 2:]

    def block(label, leaves):
        b_lines = _wrap(probe, label, f_branch, branch_w - 40)
        b_h = len(b_lines) * 44 + 30
        leaf_lines = [_wrap(probe, leaf, f_leaf, leaf_w) for leaf in leaves]
        l_h = sum(len(l) * 40 + 14 for l in leaf_lines)
        return b_lines, b_h, leaf_lines, max(b_h, l_h)

    sides = []
    for side in (right, left):
        blocks = [block(label, leaves) for label, leaves in side]
        sides.append((side, blocks, sum(b[3] for b in blocks) + gap * max(len(blocks) - 1, 0)))
    c_h = len(c_lines) * 54 + 44
    H = int(max(max(s[2] for s in sides), c_h) + 2 * margin)

    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    cy = H // 2

    for s_idx, (side, blocks, total) in enumerate(sides):
        direction = 1 if s_idx == 0 else -1
        y = (H - total) / 2
        first_index = 0 if s_idx == 0 else len(right)
        for i, (b_lines, b_h, leaf_lines, blk_h) in enumerate(blocks):
            color = _rgb(PALETTE[(first_index + i) % len(PALETTE)])
            by = y + blk_h / 2
            bw = max(d.textlength(l, font=f_branch) for l in b_lines) + 44
            bx0 = cx + direction * (c_w / 2 + branch_off) - (0 if direction > 0 else bw)
            _curve(d, (cx + direction * c_w / 2 * 0.8, cy), (bx0 if direction > 0 else bx0 + bw, by), color, 7)
            d.rounded_rectangle([bx0, by - b_h / 2, bx0 + bw, by + b_h / 2], radius=18, fill=color)
            ty = by - len(b_lines) * 44 / 2
            for line in b_lines:
                d.text((bx0 + 22, ty), line, font=f_branch, fill="white")
                ty += 44
            # leaves
            total_leaves = sum(len(l) * 40 + 14 for l in leaf_lines)
            ly = by - total_leaves / 2
            anchor_x = bx0 + bw if direction > 0 else bx0
            for lines in leaf_lines:
                lh = len(lines) * 40
                mid = ly + lh / 2
                lx = anchor_x + direction * leaf_off
                _curve(d, (anchor_x, by), (lx, mid), color, 3)
                d.ellipse([lx - 6, mid - 6, lx + 6, mid + 6], fill=color)
                tx_y = ly
                for line in lines:
                    tw = d.textlength(line, font=f_leaf)
                    tx = lx + 16 if direction > 0 else lx - 16 - tw
                    d.text((tx, tx_y), line, font=f_leaf, fill=_rgb(INK))
                    tx_y += 40
                ly += lh + 14
            y += blk_h + gap

    d.rounded_rectangle([cx - c_w / 2, cy - c_h / 2, cx + c_w / 2, cy + c_h / 2], radius=28, fill=_rgb(CENTER_FILL))
    ty = cy - len(c_lines) * 54 / 2
    for line in c_lines:
        tw = d.textlength(line, font=f_center)
        d.text((cx - tw / 2, ty), line, font=f_center, fill="white")
        ty += 54

    # crop empty horizontal margins so the map uses the available width
    bbox = Image.eval(img.convert("L"), lambda v: 255 - v).getbbox()
    if bbox:
        img = img.crop((max(bbox[0] - margin, 0), 0, min(bbox[2] + margin, W), H))
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


# ---------------------------------------------------------------- insertion

def _available_width(paragraph, doc):
    from docx.oxml.ns import qn
    from docx.shared import Emu
    tc = paragraph._p.getparent()
    if tc is not None and tc.tag == qn("w:tc"):
        tcw = tc.tcPr.find(qn("w:tcW")) if tc.tcPr is not None else None
        if tcw is not None and tcw.get(qn("w:type")) == "dxa":
            return Emu(int(tcw.get(qn("w:w"))) * 635 - 2 * 108 * 635)
    sec = doc.sections[-1]
    return Emu(sec.page_width - sec.left_margin - sec.right_margin)


def insert_mindmaps(doc, placeholders: Dict[str, Any]) -> None:
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    done = set()
    for para in paragraphs:
        if id(para._p) in done:
            continue
        done.add(id(para._p))
        m = PLACEHOLDER_RE.fullmatch(para.text.strip())
        if not m:
            continue
        value = placeholders.get(m.group(1))
        for r in list(para.runs):
            if r._r.find(qn("w:drawing")) is None:
                r._r.getparent().remove(r._r)
        if not value:
            continue
        center, branches = parse_outline(value)
        if not branches:
            para.add_run(str(value))
            continue
        try:
            png = draw_mindmap(center, branches)
        except Exception:
            logger.exception("mindmap rendering failed, falling back to the outline text")
            para.add_run(str(value))
            continue
        para.add_run().add_picture(io.BytesIO(png), width=_available_width(para, doc))


@register
class Mindmap(Renderer):
    """Mindmap: a {{mindmap_...}} placeholder becomes a mind map image drawn from an extracted outline."""
    name = "mindmap"
    family = PLACEHOLDER
    order = 20
    spec = SPEC

    def extraction_requests(self, placeholders, parse, current_metadata, force):
        return outline_requests(placeholders, parse, current_metadata, force)

    def before_substitution(self, doc, ctx):
        insert_mindmaps(doc, ctx.placeholders)
