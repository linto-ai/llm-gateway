"""Repeated rows: a table row repeated once per item of an extracted list.

Template: a row holding dotted placeholders {{actions.porteur}}, {{actions.action: hint}}... is the model
row of list `actions`. An optional {{actions: hint}} in the row describes the whole list and is removed.
Extraction: the dotted fields are requested as ONE list of objects (see extraction_requests).
Rendering: the row is duplicated per object, keeping its formatting; an empty list removes the row;
a list never extracted leaves one empty row. Top-level tables only.
"""
import re
from copy import deepcopy
from typing import Any, Dict, List, Tuple

from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .base import LISTS_KEY, PLACEHOLDER, Renderer, RendererSpec, register

SPEC = RendererSpec(
    trigger="Placeholders pointés {{liste.champ}} dans une ligne de tableau du template.",
    service_prompt="Aucune exigence propre : la sortie du service doit seulement contenir l'information (par exemple une section Actions à mener).",
    placeholder_prompt=(
        "{{liste: consigne}} dit quoi mettre dans la liste, combien d'éléments au maximum et dans quel ordre ; "
        "{{liste.champ: consigne}} dit ce que contient chaque clé. Le gateway en fait une seule demande "
        "« list of JSON objects » au prompt d'extraction, qui doit savoir renvoyer une liste d'objets."
    ),
    example="| {{actions.porteur}} | {{actions: toutes les actions, 8 au maximum}}{{actions.action: 16 mots maximum}} | {{actions.echeance}} |",
)

FIELD_RE = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\.([A-Za-z0-9_]+)\s*(?::[^}]*)?\}\}")
LIST_KEYS_FOR_STRINGS = ("value", "valeur", "texte", "text")


def list_requests(placeholders: List[str], parse, current_metadata: Dict[str, Any],
                        force: bool) -> Tuple[set, List[str]]:
    """Group dotted placeholders by list. Returns (names handled here, extraction requests to add)."""
    lists: Dict[str, Dict[str, Any]] = {}
    for placeholder in placeholders:
        info = parse(placeholder)
        if "." in info["name"]:
            name, field = info["name"].split(".", 1)
            fields = lists.setdefault(name, {"hint": None, "fields": {}})["fields"]
            if field not in fields or info["description"]:
                fields[field] = info["description"]
    for placeholder in placeholders:
        info = parse(placeholder)
        if info["name"] in lists and info["description"]:
            lists[info["name"]]["hint"] = info["description"]

    handled = set(lists)
    for placeholder in placeholders:
        name = parse(placeholder)["name"]
        if "." in name:
            handled.add(name)

    requests = []
    for name, spec in lists.items():
        if force or name not in current_metadata:
            keys = ", ".join(f"{f} ({d})" if d else f for f, d in spec["fields"].items())
            hint = f" {spec['hint']}." if spec["hint"] else ""
            requests.append(
                f"{name}: list of JSON objects, one object per item, with exactly the keys {keys}.{hint}"
                " Empty list if nothing matches."
            )
    return handled, requests


def expand_rows(doc, lists: Dict[str, Any], set_run_text, rescue) -> None:
    """Expand model rows. `lists` maps list names to extracted lists."""
    for table in doc.tables:
        for tr in list(table._tbl.findall(qn("w:tr"))):
            text = "".join(t.text or "" for t in tr.iter(qn("w:t")))
            match = FIELD_RE.search(text)
            if not match:
                continue
            name = match.group(1)
            items = lists.get(name)
            if items is None:
                items = [{}]
            if not items:
                tr.getparent().remove(tr)
                continue
            declaration = re.compile(r"\{\{\s*" + re.escape(name) + r"\s*(?::[^}]*)?\}\}")
            anchor = tr
            for item in items:
                def fill(txt, item=item):
                    def value(m):
                        if m.group(1) != name:
                            return m.group(0)
                        if isinstance(item, dict):
                            v = item.get(m.group(2), "")
                        else:
                            v = item if m.group(2) in LIST_KEYS_FOR_STRINGS else ""
                        return "" if v is None else str(v)
                    return declaration.sub("", FIELD_RE.sub(value, txt))

                row = deepcopy(tr)
                for p_el in row.iter(qn("w:p")):
                    para = Paragraph(p_el, None)
                    for run in para.runs:
                        set_run_text(run, fill(run.text))
                    rescue(para, fill)
                anchor.addnext(row)
                anchor = row
            tr.getparent().remove(tr)


@register
class RepeatedRows(Renderer):
    """Repeated rows: a table row repeated once per object of an extracted list."""
    name = "repeated_rows"
    family = PLACEHOLDER
    order = 10
    spec = SPEC

    def extraction_requests(self, placeholders, parse, current_metadata, force):
        return list_requests(placeholders, parse, current_metadata, force)

    def before_substitution(self, doc, ctx):
        expand_rows(doc, ctx.placeholders.get(LISTS_KEY) or {}, ctx.set_run_text, ctx.rescue)
