"""Helpers to build the LLM-Gateway DOCX templates (templates/default/linto-*.docx) with python-docx.

Constraints of the gateway renderer (2.6.0) that these helpers respect:
- placeholders are substituted run by run: each {{placeholder}} must sit in a single run;
- only top-level tables are visited: never nest a table holding a placeholder;
- header/footer: only paragraphs are substituted, tables there are left as is;
- {{output}} must be alone in a body paragraph (not in a table).
"""
import re
import shutil
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu

ASSETS = Path(__file__).resolve().parent.parent.parent / "templates" / "assets"
FONTS = ASSETS / "fonts"

# Brand colours: LinTO green sampled from the LinTO Studio logo, plus neutral and alert colours.
LINTO_GREEN = "1DAF92"
LINTO_GREEN_DARK = "13806B"
LINTO_GREEN_LIGHT = "E8F7F3"
ALERT_RED = "C51C42"
INK = "2D2D2D"
GREY = "6B7280"
GREY_LIGHT = "9CA3AF"
LINE = "D1D5DB"
PAPER = "F7F8FA"
WHITE = "FFFFFF"


# Successor lists from the WordprocessingML schema: children must appear in this order,
# otherwise Word may refuse the file and LibreOffice silently drops the property.
TCPR_ORDER = ("w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar",
              "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
TBLPR_ORDER = ("w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
               "w:tblLayout", "w:tblCellMar", "w:tblLook")
PPR_ORDER = ("w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr",
             "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
             "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
             "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
             "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
             "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
RPR_ORDER = ("w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike",
             "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
             "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u",
             "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
             "w:eastAsianLayout", "w:specVanish", "w:oMath")


def put(parent, child, order):
    """Insert child in parent at its schema position, replacing an existing element of the same tag."""
    tag = child.tag
    for old in parent.findall(tag):
        parent.remove(old)
    names = [qn(t) for t in order]
    idx = names.index(tag)
    parent.insert_element_before(child, *order[idx + 1:]) if hasattr(parent, "insert_element_before") else parent.append(child)
    if child.getparent() is None:
        parent.append(child)
    return child


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------- document

def new_document(margins_cm=(1.8, 1.8, 2.2, 2.4), font="Inter", size=9.5):
    """A4 document with base styles set to the brand font. margins = (left, right, top, bottom)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = [Cm(m) for m in margins_cm]
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0)
    set_style_font(doc.styles["Normal"], font, size, INK)
    for name in ("Header", "Footer"):  # built-in centre/right tabs would override ours
        doc.styles[name].paragraph_format.tab_stops.clear_all()
    pf = doc.styles["Normal"].paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    # docDefaults so that tables and lists inserted by htmldocx inherit the font too
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), font)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
    return doc


def set_style_font(style, name=None, size=None, color=None, bold=None, italic=None):
    f = style.font
    if name:
        f.name = name
        rfonts = style.element.rPr.find(qn("w:rFonts"))
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), name)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):  # theme fonts win otherwise
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = rgb(color)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def style_headings(doc, h1, h2, h3, font="Inter"):
    """h = (size, color, space_before, space_after, bottom_border_color or None)."""
    for name, (size, color, before, after, border) in (("Heading 1", h1), ("Heading 2", h2), ("Heading 3", h3)):
        st = doc.styles[name]
        set_style_font(st, font, size, color, bold=True, italic=False)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = True
        if border:
            paragraph_border(st.element.get_or_add_pPr(), bottom=(border, 6, 2))
    for name in ("List Bullet", "List Number"):
        st = doc.styles[name]
        set_style_font(st, font)
        st.paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- text

def run(paragraph, text, size=None, color=None, bold=None, italic=None, font=None, caps=False, spacing=None):
    r = paragraph.add_run(text)
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = rgb(color)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if font:
        r.font.name = font
        r._element.rPr.rFonts.set(qn("w:hAnsi"), font)
        r._element.rPr.rFonts.set(qn("w:cs"), font)
    if caps:
        r.font.all_caps = True
    if spacing is not None:  # character spacing in points
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        put(r._element.get_or_add_rPr(), sp, RPR_ORDER)
    return r


def para(container, text="", align=None, before=0, after=0, line=None, keep_next=False, **run_kw):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = line
    if align:
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if keep_next:
        pf.keep_with_next = True
    if text:
        run(p, text, **run_kw)
    return p


def first_para(cell, align=None, before=0, after=0, line=None):
    """Reuse the empty paragraph python-docx creates in every cell."""
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if line:
        pf.line_spacing = line
    if align:
        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    return p


def paragraph_border(ppr, top=None, bottom=None, left=None, right=None):
    """Each side = (hex colour, size in eighths of a point, space in points)."""
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = put(ppr, OxmlElement("w:pBdr"), PPR_ORDER)
    for side, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if spec:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(spec[1]))
            e.set(qn("w:space"), str(spec[2]))
            e.set(qn("w:color"), spec[0])
            pbdr.append(e)


def tab_stops(paragraph, stops):
    """stops = [(position_cm, 'left'|'right'|'center')]"""
    from docx.enum.text import WD_TAB_ALIGNMENT
    kinds = {"left": WD_TAB_ALIGNMENT.LEFT, "right": WD_TAB_ALIGNMENT.RIGHT, "center": WD_TAB_ALIGNMENT.CENTER}
    for pos, kind in stops:
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(pos), kinds[kind])


def field(paragraph, instr, size=None, color=None, bold=None):
    """Insert a complex field (PAGE, NUMPAGES...) as begin/instr/separate/result/end runs."""
    last = None
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("result", "1"), ("end", None)):
        r = run(paragraph, "", size=size, color=color, bold=bold)
        if kind in ("begin", "separate", "end"):
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        elif kind == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {text} "
        else:
            e = OxmlElement("w:t")
            e.text = text
        r._element.append(e)
        last = r
    return last


def picture(paragraph, path, height_cm):
    return paragraph.add_run().add_picture(str(path), height=Cm(height_cm))


# ---------------------------------------------------------------- tables

def table(container, rows, widths_cm, align="left", indent_cm=None, style=None):
    """Fixed-layout table with explicit column widths (grid + every cell)."""
    try:
        t = container.add_table(rows=rows, cols=len(widths_cm))
    except TypeError:  # header/footer containers require an explicit width
        t = container.add_table(rows, len(widths_cm), Cm(sum(widths_cm)))
    if style:
        t.style = style
    t.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "center": WD_TABLE_ALIGNMENT.CENTER}[align]
    t.autofit = False
    tblpr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    put(tblpr, layout, TBLPR_ORDER)
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(int(Cm(sum(widths_cm)).twips)))
    if indent_cm is not None:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(int(Cm(indent_cm).twips)))
        ind.set(qn("w:type"), "dxa")
        put(tblpr, ind, TBLPR_ORDER)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)
    table_borders(t, none=True)
    return t


def table_borders(t, none=False, color=LINE, size=4, inside=True, outer=True):
    tblpr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        show = not none and ((side.startswith("inside") and inside) or (not side.startswith("inside") and outer))
        e.set(qn("w:val"), "single" if show else "nil")
        if show:
            e.set(qn("w:sz"), str(size))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        b.append(e)
    put(tblpr, b, TBLPR_ORDER)


def table_cell_margins(t, top=0.1, bottom=0.1, left=0.2, right=0.2):
    tblpr = t._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(Cm(v).twips)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    put(tblpr, m, TBLPR_ORDER)


def cell_shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    put(tcpr, shd, TCPR_ORDER)


def cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (hex, eighths of pt) or 'nil'."""
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = put(tcpr, OxmlElement("w:tcBorders"), TCPR_ORDER)
    for side, spec in sides.items():
        e = OxmlElement(f"w:{side}")
        if spec == "nil":
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(spec[1]))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), spec[0])
        b.append(e)


def cell_margins(cell, top=None, bottom=None, left=None, right=None):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        if v is not None:
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(int(Cm(v).twips)))
            e.set(qn("w:type"), "dxa")
            m.append(e)
    put(tcpr, m, TCPR_ORDER)


def cell_valign(cell, where="center"):
    cell.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                               "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[where]


def row_height(row, cm, exact=False):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(Cm(cm).twips)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trpr.append(h)


def row_cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def row_header(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))


def spacer(container, pt):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pt)
    r = p.add_run("")
    r.font.size = Pt(1)
    return p


# ---------------------------------------------------------------- default table style
# htmldocx inserts markdown tables WITHOUT a table style, so Word/LibreOffice apply the
# document default table style. We turn "Normal Table" (w:default="1") into a styled grid:
# that is the only lever the 2.6.0 engine leaves to style tables coming from {{output}}.

def default_table_style(doc, border=LINE, header_fill=None, header_color=INK, cell_mar_cm=(0.08, 0.15)):
    styles = doc.styles.element
    tn = None
    for s in styles.findall(qn("w:style")):
        if s.get(qn("w:type")) == "table" and s.get(qn("w:default")) == "1":
            tn = s
    tblpr = tn.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        tn.append(tblpr)
    for child in list(tblpr):
        if child.tag in (qn("w:tblBorders"), qn("w:tblCellMar")):
            tblpr.remove(child)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), border)
        b.append(e)
    put(tblpr, b, TBLPR_ORDER)
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", cell_mar_cm[0]), ("left", cell_mar_cm[1]), ("bottom", cell_mar_cm[0]), ("right", cell_mar_cm[1])):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(int(Cm(v).twips)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    put(tblpr, m, TBLPR_ORDER)
    if header_fill:
        sp = OxmlElement("w:tblStylePr")
        sp.set(qn("w:type"), "firstRow")
        rpr = OxmlElement("w:rPr")
        col = OxmlElement("w:color")
        col.set(qn("w:val"), header_color)
        rpr.append(col)
        sp.append(rpr)
        tcpr = OxmlElement("w:tcPr")
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), header_fill)
        tcpr.append(shd)
        sp.append(tcpr)
        tn.append(sp)


# ---------------------------------------------------------------- font embedding

def _obfuscate(data: bytes, guid: str) -> bytes:
    key = bytes.fromhex(guid.strip("{}").replace("-", ""))[::-1]
    head = bytes(b ^ key[i % 16] for i, b in enumerate(data[:32]))
    return head + data[32:]


def embed_fonts(docx_path, families):
    """families = {"Inter": {"Regular": path, "Bold": path, "Italic": path, "BoldItalic": path}, ...}
    Adds obfuscated fonts (ECMA-376 embedded fonts) so Word and LibreOffice render the
    template with its real typeface even where the font is not installed (gateway pod)."""
    src = Path(docx_path)
    tmp = src.with_suffix(".tmp.docx")
    with zipfile.ZipFile(src) as zin:
        files = {n: zin.read(n) for n in zin.namelist()}
    ft = files["word/fontTable.xml"].decode()
    rels_name = "word/_rels/fontTable.xml.rels"
    rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
    fonts_xml = []
    n = 0
    for family, variants in families.items():
        embeds = []
        for variant, path in variants.items():
            n += 1
            guid = "{" + str(uuid.uuid4()).upper() + "}"
            files[f"word/fonts/font{n}.odttf"] = _obfuscate(Path(path).read_bytes(), guid)
            rels.append(f'<Relationship Id="rIdF{n}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font{n}.odttf"/>')
            embeds.append(f'<w:embed{variant} r:id="rIdF{n}" w:fontKey="{guid}"/>')
        fonts_xml.append(f'<w:font w:name="{family}"><w:charset w:val="00"/><w:family w:val="swiss"/>'
                         f'<w:pitch w:val="variable"/>{"".join(embeds)}</w:font>')
        ft = re.sub(rf'<w:font w:name="{re.escape(family)}">.*?</w:font>', "", ft, flags=re.S)
    rels.append("</Relationships>")
    if 'xmlns:r=' not in ft.split(">", 2)[1]:
        ft = ft.replace("<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ', 1)
    ft = ft.replace("</w:fonts>", "".join(fonts_xml) + "</w:fonts>")
    files["word/fontTable.xml"] = ft.encode()
    files[rels_name] = "".join(rels).encode()
    ct = files["[Content_Types].xml"].decode()
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-officedocument.obfuscatedFont"/><Default ', 1)
    files["[Content_Types].xml"] = ct.encode()
    st = files["word/settings.xml"].decode()
    if "embedTrueTypeFonts" not in st:
        st = re.sub(r"(<w:settings[^>]*>)", r'\1<w:embedTrueTypeFonts/><w:saveSubsetFonts/>', st, count=1)
    files["word/settings.xml"] = st.encode()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        order = ["[Content_Types].xml"] + [k for k in files if k != "[Content_Types].xml"]
        for k in order:
            zout.writestr(k, files[k])
    shutil.move(tmp, src)


def inter_families():
    return {
        "Inter": {"Regular": FONTS / "Inter-Regular.ttf", "Bold": FONTS / "Inter-Bold.ttf",
                  "Italic": FONTS / "Inter-Italic.ttf", "BoldItalic": FONTS / "Inter-BoldItalic.ttf"},
        "Inter Light": {"Regular": FONTS / "Inter-Light.ttf"},
        "Inter SemiBold": {"Regular": FONTS / "Inter-SemiBold.ttf"},
    }


def set_core_properties(doc, title, subject, keywords=""):
    cp = doc.core_properties
    cp.title, cp.subject, cp.author, cp.keywords = title, subject, "LinTO", keywords
    cp.comments = "Template LLM-Gateway. Placeholders {{nom: consigne}} remplis par le prompt d'extraction."


# ---------------------------------------------------------------- named table style + custom properties

def named_table_style(doc, name, border=LINE, size=4, header_fill=None, header_bold=False, header_center=True,
                      header_bottom=None, cell_mar_cm=(0.08, 0.15), font_size=None, repeat_header=True):
    """Table style meant for tables coming from {{output}}, selected through the style_table custom property."""
    from docx.enum.style import WD_STYLE_TYPE
    st = doc.styles.add_style(name, WD_STYLE_TYPE.TABLE)
    el = st.element
    tblpr = el.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        el.append(tblpr)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), border)
        b.append(e)
    put(tblpr, b, TBLPR_ORDER)
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", cell_mar_cm[0]), ("left", cell_mar_cm[1]), ("bottom", cell_mar_cm[0]), ("right", cell_mar_cm[1])):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(int(Cm(v).twips))); e.set(qn("w:type"), "dxa"); m.append(e)
    put(tblpr, m, TBLPR_ORDER)
    if font_size:
        st.font.size = Pt(font_size)
    trpr_all = OxmlElement("w:trPr"); trpr_all.append(OxmlElement("w:cantSplit"))
    tblpr.addnext(trpr_all)
    sp = OxmlElement("w:tblStylePr"); sp.set(qn("w:type"), "firstRow")
    if header_center:
        ppr = OxmlElement("w:pPr"); jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc); sp.append(ppr)
    rpr = OxmlElement("w:rPr")
    bb = OxmlElement("w:b"); bb.set(qn("w:val"), "1" if header_bold else "0"); rpr.append(bb)
    sp.append(rpr)
    if repeat_header:
        trpr = OxmlElement("w:trPr"); trpr.append(OxmlElement("w:tblHeader")); sp.append(trpr)
    tcpr = OxmlElement("w:tcPr")
    if header_bottom:
        tb = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(header_bottom[1])); e.set(qn("w:space"), "0"); e.set(qn("w:color"), header_bottom[0])
        tb.append(e); tcpr.append(tb)
    if header_fill:
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), header_fill); tcpr.append(shd)
    sp.append(tcpr)
    el.append(sp)
    return st.style_id


def set_custom_properties(docx_path, props):
    """Write docProps/custom.xml (read by the gateway: style_* today, style_table/table_widths proposed for 2.7)."""
    src = Path(docx_path)
    with zipfile.ZipFile(src) as zin:
        files = {n: zin.read(n) for n in zin.namelist()}
    items = "".join(
        f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i + 2}" name="{k}"><vt:lpwstr>{v}</vt:lpwstr></property>'
        for i, (k, v) in enumerate(props.items()))
    files["docProps/custom.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">' + items + "</Properties>").encode()
    ct = files["[Content_Types].xml"].decode()
    if "/docProps/custom.xml" not in ct:
        ct = ct.replace("</Types>", '<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>')
    files["[Content_Types].xml"] = ct.encode()
    rels = files["_rels/.rels"].decode()
    if "custom-properties" not in rels:
        rels = rels.replace("</Relationships>", '<Relationship Id="rIdCustom" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>')
    files["_rels/.rels"] = rels.encode()
    tmp = src.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for k in ["[Content_Types].xml"] + [k for k in files if k != "[Content_Types].xml"]:
            zout.writestr(k, files[k])
    shutil.move(tmp, src)


# ---------------------------------------------------------------- infographic card (shared by the service templates)

AMBER = "B45309"
CARD_GRID = [0.35, 3.95, 0.3, 3.95, 0.3, 3.95, 0.3, 3.95, 0.35]   # pad, 4 tiles with gaps, pad = 17.4 cm


class Card:
    """One top-level table on a 9-column grid; `layout` names the rows top to bottom.
    Rows named "pad" or "gap" are spacers. The gateway only substitutes placeholders in top-level
    tables, so everything of the card lives in this single table (merged cells, no nesting)."""

    def __init__(self, doc, layout, fill=PAPER):
        self.table = table(doc, len(layout), CARD_GRID)
        self.rows = {}
        for i, name in enumerate(layout):
            row = self.table.rows[i]
            row_cant_split(row)
            for c in row.cells:
                cell_shade(c, fill)
                cell_margins(c, top=0, bottom=0, left=0, right=0)
            if name in ("pad", "gap"):
                row_height(row, 0.35 if name == "pad" else 0.3, exact=True)
            else:
                self.rows[name] = row
        self.fill = fill

    def span(self, name, a, b):
        row = self.rows[name]
        return row.cells[a].merge(row.cells[b]) if a != b else row.cells[a]

    def finish(self):
        table_borders(self.table, none=True)

    def header(self, name, kicker, text_placeholder, color=LINTO_GREEN, size=11.5):
        """Coloured pill (kicker) above a lead sentence."""
        c = self.span(name, 1, 7)
        p = first_para(c, after=5)
        pill = run(p, f"   {kicker}   ", size=6.5, color=WHITE, bold=True, spacing=0.6)
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        put(pill._element.get_or_add_rPr(), shd, RPR_ORDER)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.25
        run(p, text_placeholder, size=size, color=INK)

    def tiles(self, name, tiles, accent=LINTO_GREEN, value_color=LINTO_GREEN_DARK):
        """tiles = up to 4 (label, value placeholder, value size in pt)."""
        row = self.rows[name]
        row_height(row, 1.75)
        for col, (label, value, size) in zip((1, 3, 5, 7), tiles):
            c = row.cells[col]
            cell_shade(c, WHITE)
            cell_borders(c, top=(LINE, 4), left=(LINE, 4), bottom=(accent, 18), right=(LINE, 4))
            cell_margins(c, top=0.18, bottom=0.12, left=0.3, right=0.25)
            cell_valign(c, "top")
            run(first_para(c, after=2), label, size=6.5, color=GREY, bold=True, spacing=0.5)
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            run(p, value, size=size, color=value_color if size > 12 else INK, bold=True)

    def title(self, name, text, color=INK):
        c = self.span(name, 1, 7)
        cell_margins(c, top=0.05, bottom=0.08)
        p = first_para(c)
        run(p, "●  ", size=7, color=color)
        run(p, text, size=7.5, color=color, bold=True, spacing=0.5)

    def panel(self, name, a, b, title, value, accent, size=9):
        """White panel with a coloured left rule; `value` is typically a rich multi-line placeholder."""
        cell = self.span(name, a, b)
        cell_shade(cell, WHITE)
        cell_borders(cell, top=(LINE, 4), left=(accent, 24), bottom=(LINE, 4), right=(LINE, 4))
        cell_margins(cell, top=0.2, bottom=0.2, left=0.35, right=0.3)
        cell_valign(cell, "top")
        p = first_para(cell, after=4)
        run(p, title, size=7.5, color=accent, bold=True, spacing=0.5)
        p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.25
        run(p, value, size=size, color=INK)
        return cell

    def columns(self, name, spans, labels, fill=INK, color=WHITE):
        """Dark header row of a repeated table. spans = [(a, b), ...] on the 9-column grid."""
        cells = [self.span(name, a, b) for a, b in spans]
        for c, label in zip(cells, labels):
            cell_shade(c, fill)
            cell_margins(c, top=0.08, bottom=0.08, left=0.3, right=0.2)
            run(first_para(c), label, size=6.5, color=color, bold=True, spacing=0.5)

    def repeated(self, name, spans, values):
        """Model row of a repeated table (repeated_rows renderer). values = [(placeholder, size, colour, bold)]."""
        cells = [self.span(name, a, b) for a, b in spans]
        for c, (value, size, color, bold) in zip(cells, values):
            cell_shade(c, WHITE)
            cell_borders(c, top="nil", left="nil", right="nil", bottom=(LINE, 4))
            cell_margins(c, top=0.09, bottom=0.09, left=0.3, right=0.2)
            cell_valign(c, "top")
            run(first_para(c, line=1.15), value, size=size, color=color, bold=bold)
        return cells


def brand_header_footer(section, footer_note):
    """LinTO Studio logo + organisation in the header, green rule + note + page x / y in the footer."""
    h = section.header
    h.is_linked_to_previous = False
    p = h.paragraphs[0]
    tab_stops(p, [(17.4, "right")])
    picture(p, ASSETS / "linto-studio-logo.png", 0.55)
    run(p, "\t{{organization_name}}", size=7.5, color=GREY, font="Inter Light")
    paragraph_border(p._p.get_or_add_pPr(), bottom=(LINE, 4, 6))
    f = section.footer
    f.is_linked_to_previous = False
    p = f.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    tab_stops(p, [(17.4, "right")])
    paragraph_border(p._p.get_or_add_pPr(), top=(LINTO_GREEN, 6, 6))
    run(p, footer_note, size=7, color=GREY_LIGHT, font="Inter Light")
    run(p, "\t", size=7)
    run(p, "Page ", size=7.5, color=GREY)
    field(p, "PAGE", size=7.5, color=INK, bold=True)
    run(p, " / ", size=7.5, color=GREY)
    field(p, "NUMPAGES", size=7.5, color=GREY)
    para(f, after=0)


def service_document(title_placeholder, meta_runs):
    """A4 document with brand styles, header/footer, big title and a meta line."""
    doc = new_document(margins_cm=(1.8, 1.8, 2.2, 2.2), size=10)
    doc.sections[0].header_distance = Cm(0.9)
    doc.sections[0].footer_distance = Cm(0.8)
    style_headings(doc, h1=(15, INK, 14, 6, None), h2=(12.5, INK, 12, 5, LINTO_GREEN), h3=(10.5, LINTO_GREEN_DARK, 9, 3, None))
    brand_header_footer(doc.sections[0], "Document généré avec LinTO par intelligence artificielle, à relire avant diffusion")
    para(doc, title_placeholder, after=3, size=20, color=INK, bold=True, line=1.05)
    p = para(doc, after=10)
    for i, text in enumerate(meta_runs):
        if i:
            run(p, "   ·   ", size=8.5, color=GREY_LIGHT)
        run(p, text, size=8.5, color=GREY)
    return doc


def section_label(doc, text, before=16, after=0):
    p = para(doc, before=before, after=after, keep_next=True)
    run(p, text, size=8, color=LINTO_GREEN_DARK, bold=True, spacing=0.8)
    return p
